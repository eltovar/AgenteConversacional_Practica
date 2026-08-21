#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2docx.py — Espejo Markdown → Word para la documentacion del rediseno SofIA CRM.

El Markdown vive en el repo (docs/rediseno/, versionado, para desarrollo).
El Word se genera en Desktop (legible, para el negocio).

Uso:
    python docs/rediseno/_tools/md2docx.py

Si la fuente Markdown se pierde, docx2md.py la reconstruye desde el Word.
"""
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO_DOCS = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DEST_ROOT = r"C:\Users\Salo\Desktop\Documentacion Re-estructura SofIA\Documentacion Faltante"

LAYOUT = {
    "_LEEME.md":               ("",                       "00 - LEEME (Como esta organizada esta carpeta)"),
    "00-INDICE.md":            ("",                       "00 - Indice Maestro"),
    "01-glosario.md":          ("Fase 0 - Fundamentos",   "01 - Glosario y Lenguaje Ubicuo"),
    "02-actores-y-roles.md":   ("Fase 0 - Fundamentos",   "02 - Actores y Roles"),
    "03-vision.md":            ("Fase 0 - Fundamentos",   "03 - Vision, Objetivos y No-Objetivos"),
    "14-adrs.md":              ("Fase 2 - Estructura",   "14 - Decisiones de Arquitectura (ADR)"),
    "16-antes-despues.md":     ("Fase 3 - Transicion",    "16 - Antes y Despues"),
    "17-benchmark-crms.md":    ("Fase 3 - Transicion",    "17 - Benchmark de CRMs de Referencia"),
    "wireframes/README.md":    ("Wireframes",             "Catalogo de Wireframes"),
}

ACCENT = RGBColor(0xC0, 0x39, 0x2B)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
CODE_BG = "F2F2F2"
HEAD_BG = "E8E8E8"
QUOTE_BG = "FBF6E9"


def shade(obj, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    target = obj._tc.get_or_add_tcPr() if hasattr(obj, "_tc") else obj._p.get_or_add_pPr()
    target.append(el)


def border(par, edge="left", color="C0392B", sz="18"):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), sz)
    e.set(qn("w:space"), "8" if edge == "left" else "4")
    e.set(qn("w:color"), color)
    b.append(e)
    pPr.append(b)


INLINE_RE = re.compile(
    r"(`[^`]+`)|(\*\*\*.+?\*\*\*)|(\*\*.+?\*\*)|(~~.+?~~)"
    r"|(\*[^*\n]+?\*)|(\[[^\]]+\]\([^)]+\))"
)


def add_inline(par, text):
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(9.5); r.font.color.rgb = ACCENT
        elif tok.startswith("***"):
            r = par.add_run(tok[3:-3]); r.bold = True; r.italic = True
        elif tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("~~"):
            r = par.add_run(tok[2:-2]); r.font.strike = True
        elif tok.startswith("["):
            r = par.add_run(tok[1:tok.index("]")])
            r.font.color.rgb = RGBColor(0x1F, 0x5C, 0xA8); r.underline = True
        else:
            r = par.add_run(tok[1:-1]); r.italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def add_heading(doc, level, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level <= 2 else 11)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    add_inline(p, text)
    sizes = {1: 19, 2: 15, 3: 12.5, 4: 11, 5: 10.5, 6: 10.5}
    for r in p.runs:
        r.bold = True
        r.font.size = Pt(sizes.get(level, 11))
        r.font.color.rgb = ACCENT if level <= 2 else INK
    if level == 1:
        border(p, "bottom", "C0392B", "6")


def add_code(doc, lines, lang=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    shade(p, CODE_BG)
    if lang == "mermaid":
        tag = p.add_run("DIAGRAMA \u2014 se ve dibujado en el Markdown original; "
                        "aqui va el codigo fuente\n")
    else:
        tag = p.add_run(f"[{lang}]\n" if lang else "")
    if tag.text:
        tag.font.name = "Consolas"; tag.font.size = Pt(8); tag.font.color.rgb = MUTED
    r = p.add_run("\n".join(lines))
    r.font.name = "Consolas"; r.font.size = Pt(8.5)


def add_quote(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(7)
    shade(p, QUOTE_BG)
    border(p, "left")
    for i, ln in enumerate(lines):
        if i:
            p.add_run("\n")
        add_inline(p, ln)
    for r in p.runs:
        r.font.size = Pt(10)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=1, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(ncols):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
        add_inline(p, rows[0][i] if i < len(rows[0]) else "")
        for r in p.runs:
            r.bold = True; r.font.size = Pt(9)
        shade(c, HEAD_BG)
    for row in rows[1:]:
        cells = t.add_row().cells
        for i in range(ncols):
            cells[i].text = ""
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            add_inline(p, row[i] if i < len(row) else "")
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")


def convert(md_path, docx_path, pretty_title):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().replace("\t", "    ").split("\n")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.85)
        s.top_margin = s.bottom_margin = Inches(0.8)

    p = doc.add_paragraph()
    r = p.add_run(pretty_title)
    r.bold = True; r.font.size = Pt(23); r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    rel = os.path.relpath(md_path, REPO_DOCS).replace(os.sep, "/")
    r = p.add_run(f"Redise\u00f1o SofIA CRM \u2014 Inmobiliaria Proteger   \u00b7   "
                  f"Fuente: docs/rediseno/{rel}")
    r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    p.paragraph_format.space_after = Pt(14)

    base = os.path.dirname(md_path)
    i, n, skipped_h1 = 0, len(lines), False

    while i < n:
        line, stripped = lines[i], lines[i].strip()
        if not stripped:
            i += 1; continue

        if stripped.startswith("```"):
            lang = stripped[3:].strip(); i += 1; buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            add_code(doc, buf, lang); continue

        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            border(p, "bottom", "CCCCCC", "4")
            i += 1; continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            if lvl == 1 and not skipped_h1:
                skipped_h1 = True; i += 1; continue
            add_heading(doc, lvl, txt); i += 1; continue

        mi = IMG_RE.match(stripped)
        if mi:
            path = os.path.normpath(os.path.join(base, mi.group(2)))
            if os.path.exists(path):
                try:
                    doc.add_picture(path, width=Inches(6.3))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if mi.group(1):
                        cp = doc.add_paragraph()
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cr = cp.add_run(mi.group(1))
                        cr.italic = True; cr.font.size = Pt(8.5); cr.font.color.rgb = MUTED
                except Exception:
                    pass
            i += 1; continue

        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            add_quote(doc, [b for b in buf if b.strip()]); continue

        if stripped.startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            rows = [split_row(lines[i])]; i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            add_table(doc, rows); continue

        ml = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if ml:
            depth = min(len(ml.group(1)) // 2, 3)
            ordered = ml.group(2) not in ("-", "*", "+")
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            p.paragraph_format.left_indent = Inches(0.28 + 0.25 * depth)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, ml.group(3).strip())
            for r in p.runs:
                r.font.size = Pt(10)
            i += 1; continue

        buf = [stripped]; i += 1
        while i < n:
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", ">", "|", "```", "!["))
                    or re.match(r"^(\s*)([-*+]|\d+[.)])\s+", lines[i])
                    or re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", nxt)):
                break
            buf.append(nxt); i += 1
        add_inline(doc.add_paragraph(), " ".join(buf))

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)


def main():
    made, locked = [], []
    for rel, (subdir, pretty) in LAYOUT.items():
        src = os.path.join(REPO_DOCS, rel.replace("/", os.sep))
        if not os.path.exists(src):
            print(f"  omitido (no existe): {rel}"); continue
        dst = os.path.join(DEST_ROOT, subdir, pretty + ".docx")
        try:
            convert(src, dst, pretty)
        except PermissionError:
            locked.append(os.path.join(subdir, pretty + ".docx"))
            print(f"  BLOQUEADO (abierto en Word)  {subdir}\\{pretty}.docx")
            continue
        made.append(os.path.relpath(dst, DEST_ROOT))
        print(f"  OK  {os.path.relpath(dst, DEST_ROOT)}")

    wf_src = os.path.join(REPO_DOCS, "wireframes")
    wf_dst = os.path.join(DEST_ROOT, "Wireframes", "Imagenes")
    if os.path.isdir(wf_src):
        os.makedirs(wf_dst, exist_ok=True)
        n = 0
        for f in sorted(os.listdir(wf_src)):
            if f.lower().endswith(".png"):
                shutil.copy2(os.path.join(wf_src, f), os.path.join(wf_dst, f)); n += 1
        print(f"  OK  Wireframes/Imagenes ({n} PNG)")

    print(f"\n{len(made)} documentos generados en:\n{DEST_ROOT}")
    if locked:
        print("\n" + "!" * 60)
        print("NO se pudieron actualizar (cierra Word y vuelve a ejecutar):")
        for f in locked:
            print(f"   - {f}")
        print("!" * 60)


if __name__ == "__main__":
    sys.exit(main())
